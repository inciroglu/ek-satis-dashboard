#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ek Satış Takip Paneli - build scripti

Excel'i (ek_satis.xlsx) yeniledikten sonra bu scripti çalıştırmanız yeterli:
    python build_ek_satis.py                      -> ek_satis.xlsx okunur
    python build_ek_satis.py dosya.xlsx           -> belirtilen dosya okunur
    python build_ek_satis.py --tarih 2026-10-14   -> veri tarihini elle ver

Kaynak : Excel'deki 'Konsolide Hedef' sayfası (marka bloklarındaki Plan / Fiili sütunları)
Çıktı  : index.html  (tek dosya; veri ve "Excel indir" dosyası içine gömülü)  <- GitHub'a yüklenen dosya bu
         ek_satis_takip.xlsx  (aynı Excel'in yerel kopyası, GitHub'a yüklenmez)
Kullanıcı: kullanicilar.txt (mail;sifre) varsa veri kişi bazlı şifrelenir.
Şifre  : (eski yöntem) sifre.txt dosyasına (ilk satır) şifre yazarsanız veri şifrelenir (pip install cryptography gerekir)

Not: Büyük 'Fatura Detay Liste' sayfası hiç okunmaz, bu yüzden birkaç saniyede biter.
Yeni bir ay için 'Fiili' sütunu eklendiğinde (örn. Kasım) otomatik algılanır.
Logo: İnciroğlu logosu script içine gömülüdür; isterseniz klasöre logo.svg/png koyarak değiştirebilirsiniz.
"""
import argparse
import base64
import datetime as dt
import hashlib
import json
import mimetypes
import os
import sys
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET = "Konsolide Hedef"
LEFT_COLS = 17  # sol taraftaki marka blokları (S sütunundan sonrası konsolide/parametre)

# Marka sırası ve gösterim adları (blok başlığı -> ad)
BRAND_MAP = {
    "PEUGEOT": "Peugeot",
    "CITROEN": "Citroën",
    "OPEL": "Opel",
    "FIATARJ": "Fiat, Alfa Romeo & Jeep",
    "HONDA": "Honda",
    "BMW": "BMW & MINI",
    "JAECOO": "Omoda & Jaecoo",
}
MONTHS = {"OCAK": 1, "SUBAT": 2, "MART": 3, "NISAN": 4, "MAYIS": 5, "HAZIRAN": 6,
          "TEMMUZ": 7, "AGUSTOS": 8, "EYLUL": 9, "EKIM": 10, "KASIM": 11, "ARALIK": 12}

# Çalışma günü: Cumartesi dahil, Pazar ve resmi tatiller hariç (düzenlenebilir)
HOLIDAYS = {dt.date(2026, *md) for md in [
    (1, 1), (3, 20), (3, 21), (3, 22), (4, 23), (5, 1), (5, 19),
    (5, 27), (5, 28), (5, 29), (5, 30), (7, 15), (8, 30), (10, 29)]}


def norm(v):
    """Türkçe karakter / boşluk / büyük-küçük harf farklarını yok say."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.replace("ı", "I").replace("İ", "I")
    return "".join(s.upper().split())


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def workday(d):
    return d.weekday() != 6 and d not in HOLIDAYS


def month_bounds(year, m):
    start = dt.date(year, m, 1)
    nxt = dt.date(year + (m == 12), m % 12 + 1, 1)
    return start, nxt - dt.timedelta(days=1)


def count_workdays(a, b):
    n, d = 0, a
    while d <= b:
        n += workday(d)
        d += dt.timedelta(days=1)
    return n


def parse(path, today):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"HATA: '{SHEET}' sayfası bulunamadı. Sayfalar: {wb.sheetnames}")
    rows = [list(r) for r in wb[SHEET].iter_rows(values_only=True)]
    wb.close()

    # Dönem yılı: 'Başlangıç tarihi' parametresinden
    year = today.year
    for r in rows:
        for i, c in enumerate(r):
            if norm(c) == "BASLANGICTARIHI":
                for nxt in r[i + 1:]:
                    if isinstance(nxt, (dt.datetime, dt.date)):
                        year = nxt.year
                        break

    brands, month_names, warns = {}, None, []
    for ri, r in enumerate(rows):
        t = norm(r[0] if r else None)
        if not t.endswith("SSH") or ri + 1 >= len(rows):
            continue
        key = t[:-3]
        title, sub = r[:LEFT_COLS], rows[ri + 1][:LEFT_COLS]
        mcols = [(i, MONTHS[norm(c)]) for i, c in enumerate(title) if norm(c) in MONTHS]
        if not mcols:
            continue
        end_all = next((i for i, c in enumerate(title) if norm(c) == "TOPLAM"), len(title))
        spans = []
        for j, (c0, mno) in enumerate(mcols):
            c1 = mcols[j + 1][0] if j + 1 < len(mcols) else end_all
            lab = {norm(x): c0 + k for k, x in enumerate(sub[c0:c1]) if x}
            spans.append((mno, title[c0], lab))
        if month_names is None:
            month_names = [(m, str(n)) for m, n, _ in spans]
        items, total_plan = [], {}
        for rr in rows[ri + 2:]:
            k = rr[0]
            if k is None:
                break
            if norm(k) == "TOPLAM":
                for mno, _, lab in spans:
                    if "PLAN" in lab:
                        total_plan[mno] = num(rr[lab["PLAN"]]) or 0
                break
            ms = []
            for mno, _, lab in spans:
                p = (num(rr[lab["PLAN"]]) if "PLAN" in lab else 0) or 0
                f = num(rr[lab["FIILI"]]) if "FIILI" in lab else None
                ms.append({"p": p, "f": f})
            items.append({"k": str(k).strip(), "m": ms})
        for i, (mno, _, _) in enumerate(spans):
            s = sum(it["m"][i]["p"] for it in items)
            if mno in total_plan and abs(s - total_plan[mno]) > 1:
                warns.append(f"{key} {mno}. ay: kalem toplamı {s:,.0f} ≠ sayfadaki Toplam {total_plan[mno]:,.0f}")
        brands[key] = {"key": key, "ad": BRAND_MAP.get(key, key.title()), "items": items}

    if not brands or not month_names:
        sys.exit("HATA: Marka blokları okunamadı (başlıkların 'xxx SSH' formatında olduğundan emin olun).")

    order = [k for k in BRAND_MAP if k in brands] + [k for k in brands if k not in BRAND_MAP]
    months = []
    for mno, ad in month_names:
        s, e = month_bounds(year, mno)
        durum = "kapandi" if today > e else ("gelecek" if today < s else "devam")
        tot = count_workdays(s, e)
        gec = tot if durum == "kapandi" else (0 if durum == "gelecek" else count_workdays(s, today))
        months.append({"ad": ad, "no": mno, "durum": durum, "toplamGun": tot, "gecenGun": gec,
                       "oran": gec / tot if tot else 0})
    q = (month_names[0][0] - 1) // 3 + 1
    donem = f"{year} Q{q}" if len(month_names) == 3 else f"{year} {month_names[0][1]}–{month_names[-1][1]}"
    data = {"meta": {"donem": donem, "veriTarihi": today.strftime("%d.%m.%Y"),
                     "uretim": dt.datetime.now().strftime("%d.%m.%Y %H:%M")},
            "months": months, "brands": [brands[k] for k in order]}
    return data, warns


# ----------------------------------------------------------------------------- Excel çıktısı
NAVY, BEIGE, LINE = "1F2A44", "F3EDE0", "D9CFB8"
thin = Side(style="thin", color=LINE)


def write_table(ws, title, subtitle, heads, rows, months, donem, total_label=None):
    L, nm = len(heads), len(months)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, size=9, color="6B6B6B")
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)
    for i, h in enumerate(heads):
        ws.merge_cells(start_row=4, start_column=i + 1, end_row=5, end_column=i + 1)
        c = ws.cell(4, i + 1, h)
        c.font, c.fill, c.alignment = hf, fill, Alignment(horizontal="left", vertical="center")
        ws.cell(5, i + 1).fill = fill
    groups = [(m["ad"]) for m in months] + [f"{donem} Toplam"]
    for g, name in enumerate(groups):
        c0 = L + 1 + 3 * g
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
        c = ws.cell(4, c0, name)
        c.font, c.fill, c.alignment = hf, fill, Alignment(horizontal="center")
        for k, sub in enumerate(["Plan", "Fiili", "Gerç. %"]):
            s = ws.cell(5, c0 + k, sub)
            s.font, s.fill, s.alignment = hf, PatternFill("solid", fgColor="34425F"), Alignment(horizontal="center")
            ws.cell(4, c0 + k).fill = fill

    def pc(r, g, k):
        return f"{get_column_letter(L + 1 + 3 * g + k)}{r}"

    def fill_formulas(r):
        for g in range(nm):
            p, f = pc(r, g, 0), pc(r, g, 1)
            ws[pc(r, g, 2)] = f'=IF(OR({p}=0,{f}=""),"",{f}/{p})'
        ps = ",".join(pc(r, g, 0) for g in range(nm))
        fs = ",".join(pc(r, g, 1) for g in range(nm))
        q = nm
        ws[pc(r, q, 0)] = f"=SUM({ps})"
        ws[pc(r, q, 1)] = f'=IF(COUNT({fs})=0,"",SUM({fs}))'
        ws[pc(r, q, 2)] = f'=IF(OR({pc(r, q, 0)}=0,{pc(r, q, 1)}=""),"",{pc(r, q, 1)}/{pc(r, q, 0)})'

    r = 6
    first = r
    for labels, vals in rows:
        for i, lab in enumerate(labels):
            ws.cell(r, i + 1, lab)
        for g, (p, f) in enumerate(vals):
            ws[pc(r, g, 0)] = p
            ws[pc(r, g, 1)] = f
        fill_formulas(r)
        r += 1
    last = r - 1
    if total_label:
        ws.cell(r, 1, total_label)
        for g in range(nm):
            for k in (0, 1):
                col = get_column_letter(L + 1 + 3 * g + k)
                rng = f"{col}{first}:{col}{last}"
                ws[pc(r, g, k)] = f"=SUM({rng})" if k == 0 else f'=IF(COUNT({rng})=0,"",SUM({rng}))'
        fill_formulas(r)
        for c in range(1, L + 3 * (nm + 1) + 1):
            cell = ws.cell(r, c)
            cell.font, cell.fill = Font(bold=True, color=NAVY), PatternFill("solid", fgColor=BEIGE)
    for rr in range(first, r + 1):
        for c in range(L + 1, L + 3 * (nm + 1) + 1):
            cell = ws.cell(rr, c)
            cell.number_format = "0.0%" if (c - L) % 3 == 0 else "#,##0"
            cell.border = Border(bottom=thin)
        for c in range(1, L + 1):
            ws.cell(rr, c).border = Border(bottom=thin)
    for i in range(L):
        ws.column_dimensions[get_column_letter(i + 1)].width = 28 if i == 0 else 18
    for c in range(L + 1, L + 3 * (nm + 1) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = ws.cell(6, L + 1)
    return first, last


def write_xlsx(data, path):
    ms, donem = data["months"], data["meta"]["donem"]
    sub = f"Veri tarihi: {data['meta']['veriTarihi']} · Tutarlar TL · Kaynak: Konsolide Hedef sayfası"
    nm = len(ms)

    def merge(lists):
        out = [[0, None] for _ in range(nm)]
        for lst in lists:
            for i, m in enumerate(lst):
                out[i][0] += m["p"]
                if m["f"] is not None:
                    out[i][1] = (out[i][1] or 0) + m["f"]
        return [tuple(x) for x in out]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marka Özeti"
    rows = [([b["ad"]], merge([it["m"] for it in b["items"]])) for b in data["brands"]]
    write_table(ws, f"Ek Satış Takip · {donem} · Marka Özeti", sub, ["Marka"], rows, ms, donem, "İNCİROĞLU TOPLAM")

    ws = wb.create_sheet("Kalem Özeti")
    kalems = []
    for b in data["brands"]:
        for it in b["items"]:
            if it["k"] not in kalems:
                kalems.append(it["k"])
    rows = [([k], merge([it["m"] for b in data["brands"] for it in b["items"] if it["k"] == k])) for k in kalems]
    write_table(ws, f"Ek Satış Takip · {donem} · Kalem Özeti", sub, ["Kalem"], rows, ms, donem, "İNCİROĞLU TOPLAM")

    ws = wb.create_sheet("Marka-Kalem Detay")
    rows = [([b["ad"], it["k"]], [(m["p"], m["f"]) for m in it["m"]])
            for b in data["brands"] for it in b["items"]]
    first, last = write_table(ws, f"Ek Satış Takip · {donem} · Marka × Kalem Detay", sub, ["Marka", "Kalem"],
                              rows, ms, donem)
    ws.auto_filter.ref = f"A5:{get_column_letter(2 + 3 * (nm + 1))}{last}"
    wb.save(path)


# ----------------------------------------------------------------------------- HTML şablonu
LOGO_SVG = r'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="60 383 490 76" role="img" aria-label="İnciroğlu Otomotiv"><path fill="currentColor" d="M 61.92 398.42 L 71.32 398.42 L 71.32 443.6 L 61.92 443.6 Z M 66.68 394.76 C 64.85 394.76 63.54 394.38 62.75 393.6 C 61.95 392.83 61.55 391.55 61.55 389.75 C 61.55 387.96 61.95 386.68 62.75 385.91 C 63.54 385.13 64.85 384.75 66.68 384.75 C 68.48 384.75 69.79 385.14 70.62 385.94 C 71.46 386.73 71.88 388 71.88 389.75 C 71.88 391.5 71.46 392.78 70.62 393.57 C 69.79 394.36 68.48 394.76 66.68 394.76  M 78.53 410.63 L 86.28 410.63 L 86.83 414.23 L 87.14 414.23 C 88.48 413.46 89.68 412.8 90.74 412.25 C 91.8 411.7 92.79 411.27 93.7 410.94 C 94.62 410.61 95.49 410.38 96.33 410.24 C 97.16 410.09 98 410.02 98.86 410.02 C 101.42 410.02 103.39 410.8 104.78 412.34 C 106.16 413.89 106.86 416.09 106.86 418.94 L 106.86 443.6 L 97.64 443.6 L 97.64 422.05 C 97.64 420.5 97.36 419.41 96.81 418.75 C 96.27 418.11 95.46 417.78 94.4 417.78 C 93.46 417.78 92.42 417.97 91.26 418.36 C 90.1 418.75 88.93 419.26 87.75 419.91 L 87.75 443.6 L 78.53 443.6 Z M 78.53 410.63  M 127.92 444.15 C 122.59 444.15 118.6 442.78 115.95 440.03 C 113.31 437.29 111.98 432.98 111.98 427.12 C 111.98 421.3 113.31 417 115.95 414.23 C 118.6 411.47 122.59 410.09 127.92 410.09 C 129.22 410.09 130.52 410.16 131.8 410.3 C 133.08 410.44 134.25 410.68 135.31 411 L 135.31 418.27 C 134.12 418.06 133 417.9 131.92 417.78 C 130.84 417.66 129.83 417.59 128.89 417.59 C 126.17 417.59 124.23 418.32 123.07 419.76 C 121.91 421.21 121.32 423.64 121.32 427.05 C 121.32 430.48 121.91 432.94 123.07 434.45 C 124.23 435.95 126.15 436.7 128.84 436.7 C 130.02 436.7 131.11 436.65 132.13 436.55 C 133.15 436.45 134.21 436.27 135.31 436.03 L 135.31 443.18 C 134.37 443.5 133.24 443.74 131.92 443.91 C 130.59 444.07 129.26 444.15 127.92 444.15  M 140.68 410.63 L 149.9 410.63 L 149.9 443.6 L 140.68 443.6 Z M 145.32 405.63 C 143.49 405.63 142.18 405.23 141.38 404.44 C 140.59 403.64 140.19 402.37 140.19 400.62 C 140.19 398.87 140.59 397.6 141.38 396.8 C 142.18 396.01 143.49 395.61 145.32 395.61 C 147.15 395.61 148.46 396.02 149.26 396.84 C 150.05 397.65 150.45 398.91 150.45 400.62 C 150.45 402.33 150.05 403.59 149.26 404.41 C 148.46 405.22 147.15 405.63 145.32 405.63  M 156.49 410.63 L 164.24 410.63 L 164.79 415.34 L 165.16 415.34 C 165.57 414.93 166.11 414.48 166.81 413.99 C 167.5 413.5 168.27 413.02 169.13 412.53 C 169.98 412.04 170.86 411.57 171.75 411.12 C 172.65 410.67 173.5 410.33 174.32 410.09 L 176.52 410.09 L 176.52 418.45 L 172.98 418.45 C 171.67 418.65 170.32 419.04 168.91 419.61 C 167.51 420.18 166.44 420.77 165.71 421.38 L 165.71 443.6 L 156.49 443.6 Z M 156.49 410.63  M 193.67 437.13 C 194.81 437.13 195.77 436.97 196.54 436.64 C 197.31 436.32 197.95 435.78 198.43 435.02 C 198.92 434.27 199.27 433.25 199.47 431.97 C 199.68 430.69 199.78 429.07 199.78 427.12 C 199.78 425.16 199.68 423.55 199.47 422.26 C 199.27 420.98 198.92 419.96 198.43 419.21 C 197.95 418.46 197.31 417.92 196.54 417.59 C 195.77 417.27 194.81 417.11 193.67 417.11 C 192.53 417.11 191.58 417.27 190.8 417.59 C 190.03 417.92 189.41 418.46 188.94 419.21 C 188.47 419.96 188.14 420.98 187.93 422.26 C 187.73 423.55 187.63 425.16 187.63 427.12 C 187.63 429.07 187.73 430.69 187.93 431.97 C 188.14 433.25 188.47 434.27 188.94 435.02 C 189.41 435.78 190.03 436.32 190.8 436.64 C 191.58 436.97 192.53 437.13 193.67 437.13 M 193.67 444.21 C 188.66 444.21 184.86 442.77 182.25 439.88 C 179.65 436.99 178.35 432.73 178.35 427.12 C 178.35 421.5 179.65 417.25 182.25 414.36 C 184.86 411.47 188.66 410.02 193.67 410.02 C 198.72 410.02 202.54 411.47 205.12 414.36 C 207.7 417.25 209 421.5 209 427.12 C 209 432.73 207.7 436.99 205.12 439.88 C 202.54 442.77 198.72 444.21 193.67 444.21  M 227.74 436.52 C 229.57 436.52 230.97 435.87 231.95 434.57 C 232.93 433.27 233.41 431.07 233.41 427.97 L 233.41 416.98 C 232.15 416.58 230.67 416.38 228.96 416.38 C 227.74 416.38 226.7 416.55 225.84 416.89 C 224.99 417.24 224.31 417.81 223.8 418.6 C 223.29 419.39 222.91 420.47 222.67 421.84 C 222.43 423.2 222.3 424.88 222.3 426.87 C 222.3 428.66 222.39 430.17 222.58 431.39 C 222.76 432.61 223.07 433.61 223.5 434.38 C 223.92 435.16 224.48 435.7 225.17 436.03 C 225.87 436.36 226.72 436.52 227.74 436.52 M 229.08 406.12 C 225.25 406.12 222.43 405.25 220.59 403.52 C 218.76 401.79 217.79 399.38 217.66 396.29 L 223.95 396.29 C 223.99 399.62 225.7 401.29 229.08 401.29 C 232.42 401.29 234.11 399.62 234.15 396.29 L 240.5 396.29 C 240.34 399.42 239.34 401.84 237.51 403.55 C 235.68 405.26 232.87 406.12 229.08 406.12 M 226.03 457.03 C 224.24 457.03 222.41 456.91 220.56 456.67 C 218.71 456.42 216.93 456.08 215.22 455.63 L 215.83 448.85 C 219.09 449.67 222.1 450.07 224.87 450.07 C 227.8 450.07 230 449.39 231.49 448.03 C 232.98 446.66 233.74 444.29 233.78 440.91 L 233.78 438.78 L 233.41 438.78 C 231.46 441.46 228.59 442.81 224.8 442.81 C 220.86 442.81 217.9 441.59 215.92 439.14 C 213.95 436.7 212.96 432.92 212.96 427.79 C 212.96 421.93 214.37 417.51 217.18 414.54 C 219.98 411.57 223.99 410.09 229.2 410.09 C 231.07 410.09 233.05 410.24 235.12 410.54 C 237.2 410.85 239.7 411.39 242.63 412.16 L 242.63 439.33 C 242.63 442.42 242.26 445.07 241.5 447.3 C 240.75 449.51 239.65 451.34 238.21 452.79 C 236.76 454.23 235.01 455.3 232.96 456 C 230.9 456.69 228.59 457.03 226.03 457.03  M 249.23 394.27 L 258.45 394.27 L 258.45 443.6 L 249.23 443.6 Z M 249.23 394.27  M 273.59 444.21 C 270.82 444.21 268.67 443.43 267.14 441.86 C 265.62 440.3 264.86 438.11 264.86 435.3 L 264.86 410.63 L 274.07 410.63 L 274.07 432.43 C 274.07 433.93 274.35 435.02 274.9 435.66 C 275.45 436.32 276.34 436.64 277.55 436.64 C 278.49 436.64 279.49 436.45 280.55 436.06 C 281.61 435.68 282.76 435.12 284.02 434.38 L 284.02 410.63 L 293.25 410.63 L 293.25 443.6 L 285.49 443.6 L 284.94 440.06 L 284.7 440.06 C 283.39 440.83 282.27 441.48 281.31 442.02 C 280.35 442.54 279.46 442.97 278.62 443.3 C 277.79 443.62 276.98 443.86 276.18 444 C 275.39 444.14 274.52 444.21 273.59 444.21  M 328.11 439.94 C 332.42 439.94 335.55 438.43 337.5 435.42 C 339.46 432.41 340.44 427.62 340.44 421.07 C 340.44 414.52 339.44 409.71 337.45 406.64 C 335.45 403.56 332.34 402.03 328.11 402.03 C 323.87 402.03 320.77 403.55 318.8 406.61 C 316.82 409.66 315.83 414.48 315.83 421.07 C 315.83 427.62 316.82 432.41 318.8 435.42 C 320.77 438.43 323.87 439.94 328.11 439.94 M 328.11 444.21 C 322.41 444.21 318.11 442.29 315.22 438.44 C 312.33 434.6 310.89 428.81 310.89 421.07 C 310.89 413.26 312.33 407.43 315.22 403.58 C 318.11 399.73 322.41 397.81 328.11 397.81 C 333.84 397.81 338.17 399.73 341.08 403.58 C 343.99 407.43 345.45 413.26 345.45 421.07 C 345.45 428.81 343.99 434.6 341.08 438.44 C 338.17 442.29 333.84 444.21 328.11 444.21  M 364.07 443.97 C 358.2 443.97 355.27 440.98 355.27 434.99 L 355.27 414.91 L 349.11 414.91 L 349.11 410.63 L 355.52 410.63 L 356.19 402.39 L 360.04 401.78 L 360.04 410.63 L 368.83 410.63 L 368.83 414.91 L 360.04 414.91 L 360.04 434.08 C 360.04 438.03 361.77 440 365.23 440 C 365.71 440 366.28 439.97 366.93 439.91 C 367.59 439.85 368.22 439.71 368.83 439.51 L 368.83 443.42 C 367.16 443.79 365.57 443.97 364.07 443.97  M 385.8 440.18 C 388.77 440.18 390.94 439.14 392.3 437.07 C 393.66 434.99 394.34 431.68 394.34 427.12 C 394.34 422.56 393.65 419.24 392.27 417.17 C 390.89 415.09 388.73 414.05 385.8 414.05 C 382.82 414.05 380.66 415.09 379.3 417.17 C 377.93 419.24 377.25 422.56 377.25 427.12 C 377.25 431.68 377.93 434.99 379.3 437.07 C 380.66 439.14 382.82 440.18 385.8 440.18 M 385.8 444.15 C 381.36 444.15 378 442.71 375.72 439.82 C 373.45 436.93 372.3 432.7 372.3 427.12 C 372.3 421.54 373.46 417.3 375.75 414.39 C 378.05 411.48 381.4 410.02 385.8 410.02 C 390.2 410.02 393.53 411.48 395.81 414.39 C 398.09 417.3 399.23 421.54 399.23 427.12 C 399.23 432.7 398.09 436.93 395.81 439.82 C 393.53 442.71 390.2 444.15 385.8 444.15  M 405.88 410.63 L 409.79 410.63 L 410.16 414.3 L 410.4 414.3 C 412.6 412.83 414.56 411.75 416.29 411.06 C 418.02 410.37 419.68 410.02 421.27 410.02 C 424.93 410.02 427.23 411.47 428.17 414.36 L 428.47 414.36 C 430.67 412.81 432.64 411.7 434.39 411.03 C 436.14 410.36 437.84 410.02 439.46 410.02 C 442.07 410.02 443.98 410.71 445.2 412.1 C 446.42 413.48 447.04 415.72 447.04 418.81 L 447.04 443.6 L 442.27 443.6 L 442.27 420.16 C 442.27 418.08 441.91 416.59 441.17 415.67 C 440.44 414.75 439.3 414.3 437.75 414.3 C 436.45 414.3 435.06 414.58 433.57 415.15 C 432.09 415.72 430.51 416.54 428.84 417.59 L 428.84 443.6 L 424.08 443.6 L 424.08 420.16 C 424.08 418.08 423.71 416.59 422.98 415.67 C 422.25 414.75 421.11 414.3 419.56 414.3 C 418.26 414.3 416.85 414.59 415.35 415.18 C 413.84 415.77 412.25 416.6 410.59 417.66 L 410.59 443.6 L 405.88 443.6 Z M 405.88 410.63  M 466.93 440.18 C 469.91 440.18 472.07 439.14 473.44 437.07 C 474.8 434.99 475.48 431.68 475.48 427.12 C 475.48 422.56 474.79 419.24 473.41 417.17 C 472.02 415.09 469.87 414.05 466.93 414.05 C 463.96 414.05 461.8 415.09 460.43 417.17 C 459.07 419.24 458.39 422.56 458.39 427.12 C 458.39 431.68 459.07 434.99 460.43 437.07 C 461.8 439.14 463.96 440.18 466.93 440.18 M 466.93 444.15 C 462.5 444.15 459.14 442.71 456.86 439.82 C 454.58 436.93 453.44 432.7 453.44 427.12 C 453.44 421.54 454.59 417.3 456.89 414.39 C 459.19 411.48 462.54 410.02 466.93 410.02 C 471.33 410.02 474.67 411.48 476.95 414.39 C 479.23 417.3 480.37 421.54 480.37 427.12 C 480.37 432.7 479.23 436.93 476.95 439.82 C 474.67 442.71 471.33 444.15 466.93 444.15  M 498.2 443.97 C 492.33 443.97 489.4 440.98 489.4 434.99 L 489.4 414.91 L 483.24 414.91 L 483.24 410.63 L 489.65 410.63 L 490.32 402.39 L 494.16 401.78 L 494.16 410.63 L 502.96 410.63 L 502.96 414.91 L 494.16 414.91 L 494.16 434.08 C 494.16 438.03 495.89 440 499.36 440 C 499.84 440 500.41 439.97 501.06 439.91 C 501.71 439.85 502.35 439.71 502.96 439.51 L 502.96 443.42 C 501.29 443.79 499.7 443.97 498.2 443.97  M 508.09 410.63 L 512.85 410.63 L 512.85 443.6 L 508.09 443.6 Z M 510.46 403.43 C 508.43 403.43 507.41 402.41 507.41 400.38 C 507.41 398.38 508.43 397.39 510.46 397.39 C 512.54 397.39 513.58 398.38 513.58 400.38 C 513.58 402.41 512.54 403.43 510.46 403.43  M 517.85 411.06 L 517.85 410.63 L 522.86 410.63 L 530.79 435.48 L 531.83 439.33 L 532.32 439.33 L 533.36 435.48 L 541.36 410.63 L 546.24 410.63 L 546.24 411.06 L 535.13 443.6 L 528.96 443.6 Z M 517.85 411.06 "/></svg>'''


HTML = r'''<!DOCTYPE html>
<html lang="tr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex,nofollow">
<title>Ek Satış Takip Paneli</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
<link rel="icon" href="data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 64 64'%3E%3Crect width='64' height='64' rx='12' fill='%23122f59'/%3E%3Ctext x='32' y='47' font-family='Segoe UI,Arial,sans-serif' font-size='40' font-weight='700' fill='white' text-anchor='middle'%3Eİ%3C/text%3E%3C/svg%3E">
<style>
:root{--navy:#122f59;--navy2:#183e6d;--ink:#192c44;--muted:#64758a;--line:#e1e7ed;--bg:#f4f6f9;--green:#087565;--red:#b44948;--gold:#b78d4c;--soft:#eaf0f7;--amber:#8a601e;--plan:#b8c7dc}
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:"Segoe UI","Inter",system-ui,-apple-system,"Helvetica Neue",Arial,sans-serif;font-size:14px;line-height:1.5;color:var(--ink);background:var(--bg);-webkit-font-smoothing:antialiased;text-rendering:optimizeLegibility}
button,a.btn{font:inherit;color:var(--ink);cursor:pointer;border:1px solid #cfd9e5;background:#fff;border-radius:7px;padding:8px 14px;font-weight:600;text-decoration:none;display:inline-block;transition:background .15s,border-color .15s}
button:hover,a.btn:hover{background:#f1f5fa}
button.pri,a.btn.pri{background:var(--navy);border-color:var(--navy);color:#fff}
button.pri:hover{background:var(--navy2)}
:focus-visible{outline:2px solid var(--gold);outline-offset:2px}
table,.v,.chip,td,.cv{font-variant-numeric:tabular-nums lining-nums;font-feature-settings:"tnum" 1,"lnum" 1}

/* üst bar */
.top{background:#fff;border-bottom:1px solid var(--line)}
.top .in{max-width:1320px;margin:0 auto;display:flex;align-items:center;gap:12px;padding:14px 28px;flex-wrap:wrap}
.logo{width:212px;max-width:50%;color:#231f20}.logo svg{display:block;width:100%;height:auto}
.wm{font-size:21px;font-weight:700;letter-spacing:-.6px;color:var(--navy)}
.acts{margin-left:auto;display:flex;gap:8px}
main{max-width:1320px;margin:0 auto;padding:6px 28px 36px}

/* başlık */
.head{display:flex;align-items:flex-end;gap:24px;margin:20px 0 18px;flex-wrap:wrap}
.kick{font-size:11px;font-weight:600;letter-spacing:.07em;text-transform:uppercase;color:#64809f;margin-bottom:4px}
.head h1{font-size:30px;letter-spacing:-.02em;color:var(--navy);line-height:1.1;font-weight:700}
.head .sub{color:var(--muted);font-size:13px;margin-top:5px}
.hr{margin-left:auto;display:flex;align-items:center;gap:18px}
.stamp{width:110px;height:110px;border-radius:50%;flex:none;position:relative;color:#fff;text-align:center;display:flex;flex-direction:column;align-items:center;justify-content:center;line-height:1.1;
 background:radial-gradient(circle at 30% 25%,#12987f,#087565 60%,#06604f);box-shadow:0 6px 16px rgba(8,117,101,.28),inset 0 0 0 3px rgba(255,255,255,.22)}
.stamp::before{content:"";position:absolute;inset:-6px;border-radius:50%;border:2px solid rgba(8,117,101,.35);animation:ring 2.6s ease-out infinite}
.stamp small{font-size:8px;font-weight:700;letter-spacing:.1em;opacity:.88;line-height:1.25}
.stamp b{font-size:30px;font-weight:700;letter-spacing:-.03em;margin:2px 0 1px}
.stamp span{font-size:11.5px;font-weight:600;opacity:.95}
@keyframes ring{0%{transform:scale(.94);opacity:.9}100%{transform:scale(1.18);opacity:0}}
.days{background:#fff;border:1px solid var(--line);border-radius:9px;padding:10px 18px 9px;width:270px}
.days .l{font-size:11.5px;color:var(--muted);font-weight:600;display:flex;justify-content:space-between;gap:8px;white-space:nowrap}
.days .l b{color:var(--navy)}
.days .bar{position:relative;height:10px;background:var(--soft);border-radius:5px;margin:24px 0 5px}
.days .bar i.fb{position:absolute;left:0;top:0;display:block;height:100%;width:0;border-radius:5px;transition:width 1.4s cubic-bezier(.2,.8,.2,1);
 background:repeating-linear-gradient(115deg,rgba(255,255,255,.28) 0 8px,rgba(255,255,255,0) 8px 16px),linear-gradient(90deg,#183e6d,#087565 70%,#12b08f);background-size:32px 100%,100% 100%;animation:stripes 1.2s linear infinite}
@keyframes stripes{to{background-position:32px 0,0 0}}
.days .cur{position:absolute;top:50%;left:0;width:20px;height:20px;margin:-10px 0 0 -10px;border-radius:50%;background:#fff;border:4px solid #087565;box-shadow:0 2px 8px rgba(8,117,101,.45);transition:left 1.4s cubic-bezier(.2,.8,.2,1);z-index:2}
.days .cur::after{content:"";position:absolute;inset:-9px;border-radius:50%;border:2px solid rgba(8,117,101,.55);animation:ping 1.8s ease-out infinite}
@keyframes ping{0%{transform:scale(.6);opacity:1}100%{transform:scale(1.4);opacity:0}}
.days .cur span{position:absolute;bottom:25px;left:50%;transform:translateX(-50%);background:var(--navy);color:#fff;font-size:10.5px;font-weight:700;padding:2px 8px;border-radius:5px;white-space:nowrap}
.days .cur span::after{content:"";position:absolute;top:100%;left:50%;margin-left:-4px;border:4px solid transparent;border-top-color:var(--navy)}
.days .ends{display:flex;justify-content:space-between;font-size:10.5px;color:#8a98aa;font-weight:600}

/* KPI kartları */
.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin-bottom:22px}
.card{background:#fff;border:1px solid var(--line);border-top:3px solid var(--navy);border-radius:9px;padding:14px 18px;animation:up .5s ease both;transition:transform .2s,box-shadow .2s}
.card:nth-child(2){animation-delay:.06s}.card:nth-child(3){animation-delay:.12s}.card:nth-child(4){animation-delay:.18s}
.card:hover{transform:translateY(-2px);box-shadow:0 8px 20px rgba(18,47,89,.12)}
.card .t{font-size:11.5px;font-weight:600;letter-spacing:.07em;text-transform:uppercase;color:#47607f}
.card .v{font-size:27px;font-weight:700;letter-spacing:-.02em;color:var(--navy);margin-top:6px;line-height:1.1}
.card .s{font-size:12px;color:var(--muted);margin-top:4px}
.card.g{border-top-color:var(--green)}.card.g .v{color:var(--green)}
.card.a{border-top-color:#c58a2b}.card.a .v{color:var(--amber)}
.card.r{border-top-color:var(--red)}.card.r .v{color:var(--red)}
.card.n{border-top-color:#9fb3cc}

/* sekmeler */
.tabs{display:flex;gap:4px;border-bottom:1px solid var(--line);margin-bottom:18px;flex-wrap:wrap}
.tab{border:0;border-bottom:3px solid transparent;border-radius:0;background:none;padding:11px 20px;font-size:15px;font-weight:600;color:var(--muted);margin-bottom:-1px;transition:color .2s,border-color .2s,background .2s}
.tab:hover{color:var(--navy);background:#eef2f7}
.tab.on{color:var(--navy);font-weight:700;border-bottom-color:var(--gold);background:none}
.panel{display:none}.panel.on{display:block;animation:up .45s ease both}
@keyframes up{from{opacity:0;transform:translateY(10px)}to{opacity:1;transform:none}}
@keyframes fade{from{opacity:0}to{opacity:1}}

/* kutular */
.box{background:#fff;border:1px solid var(--line);border-radius:9px;margin-bottom:18px;overflow:hidden}
.boxhead{display:flex;align-items:center;gap:12px;padding:12px 18px;border-bottom:1px solid var(--line);flex-wrap:wrap}
.boxhead h2{font-size:15px;color:var(--navy);font-weight:700}
.boxhead span.s{font-size:12px;color:var(--muted)}
.sg{margin-left:auto;display:flex;gap:4px;flex-wrap:wrap}
.sg button{padding:4px 12px;font-size:12px;border-radius:14px}
.sg button.on{background:var(--navy);color:#fff;border-color:var(--navy)}
.pad{padding:12px 18px 14px}
.scroll{overflow-x:auto}
.bchips{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:18px}
.bchip{padding:9px 16px;border-radius:9px;display:flex;flex-direction:column;align-items:flex-start;line-height:1.25;transition:all .2s}
.bchip small{font-weight:500;font-size:11.5px;color:var(--muted)}
.bchip:hover{border-color:var(--navy);transform:translateY(-1px)}
.bchip.on{background:var(--navy);border-color:var(--navy);color:#fff}.bchip.on small{color:#c4d0e0}

/* grafik */
.lg{display:flex;flex-wrap:wrap;gap:16px;font-size:12px;color:var(--muted);padding:0 0 6px}
.lg i{display:inline-block;width:12px;height:12px;border-radius:3px;margin-right:6px;vertical-align:-2px}
.crow{display:grid;grid-template-columns:minmax(110px,190px) 1fr 80px;gap:14px;align-items:center;padding:9px 0;border-top:1px solid var(--line)}
.crow:first-of-type{border-top:0}
.cn{font-weight:600;font-size:13.5px;color:var(--navy)}
.ct{position:relative;height:17px;margin:2px 0}
.cf{position:absolute;left:0;top:0;height:100%;width:0;border-radius:0 4px 4px 0;transition:width 1.1s cubic-bezier(.2,.8,.2,1)}
.cf.pl{background:var(--plan)}.cf.fi{background:var(--navy)}.cf.nz{min-width:3px}
.cv{position:absolute;left:6px;top:0;line-height:17px;font-size:12.5px;color:var(--muted);white-space:nowrap;transition:left 1.1s cubic-bezier(.2,.8,.2,1)}
.cv.b{color:var(--navy);font-weight:700}
.cp{text-align:right}
.chip{display:inline-block;font-weight:700;padding:3px 8px;border-radius:5px;font-size:12.5px;white-space:nowrap}
.chip.g{background:#e8f4ef;color:var(--green);box-shadow:inset 0 0 0 1px #bfe0d3}
.chip.a{background:#faf1df;color:var(--amber);box-shadow:inset 0 0 0 1px #ecd7ad}
.chip.r{background:#faeeee;color:var(--red);box-shadow:inset 0 0 0 1px #ecc8c7}
.chip.n{background:#f1f4f7;color:var(--muted);font-weight:500;box-shadow:inset 0 0 0 1px var(--line)}

/* tablo */
table{width:100%;border-collapse:collapse;font-size:13.5px}
th,td{padding:9px 12px;text-align:right;white-space:nowrap}
th:first-child,td:first-child{text-align:left}
tr.sub th:first-child{text-align:right}
thead th{font-size:11.5px;color:#5d6f84;font-weight:600;background:#f7f9fb}
thead tr.per th{font-size:12.5px;color:var(--navy);text-align:center;background:#fff;border-bottom:1px solid var(--line);border-left:1px solid var(--line);font-weight:700}
thead tr.per th:first-child{border-left:0}
thead tr.per th.cur{border-bottom:2px solid var(--gold)}
thead tr.per th.ytd{border-bottom:2px solid var(--navy)}
thead tr.per small{font-weight:500;color:var(--muted);font-size:11px}
tbody tr{border-top:1px solid var(--line);transition:background .15s;animation:fade .5s both}
tbody tr:hover{background:#fafcfe}
td:first-child{font-weight:600}
td.sep{border-left:1px solid var(--line)}
td.na{color:#a9b4c1}
td.pp{font-weight:650}
td.pp.g{color:var(--green);background:#e6f3ef}td.pp.a{color:var(--amber);background:#fbf3e3}td.pp.r{color:var(--red);background:#f9ebea}
td.qq{background:#f7f9fb}
tfoot td{font-weight:700;border-top:2px solid var(--navy);background:#fafbfc;color:var(--navy)}
tfoot td.pp{background:none}
.dot{display:inline-block;width:9px;height:9px;border-radius:50%;margin-right:8px}
.note{font-size:12px;color:var(--muted);margin:2px 2px 0;line-height:1.7}
.note .chip{font-size:11px;padding:1px 7px;margin:0 2px}

#printArea{display:none}
.panel.noanim{animation:none}
.rkw{min-width:780px}
.rkrow{display:grid;grid-template-columns:64px minmax(170px,1.8fr) repeat(4,minmax(100px,1fr));align-items:center;padding:13px 18px;border-top:1px solid var(--line);background:#fff;position:relative;font-variant-numeric:tabular-nums}
.rkrow>div:nth-child(n+3){text-align:right}.rkrow>div:first-child{text-align:center}
.rkrow .nm{font-weight:650;color:var(--navy);font-size:14.5px}
.rkrow.rkhead{background:#f7f9fb;font-size:11.5px;font-weight:600;color:#5d6f84;border-top:0;padding:9px 18px}
.rkrow.first{background:#fdf9f1}
.rkrow.tot{border-top:2px solid var(--navy);background:#fafbfc;font-weight:700;color:var(--navy)}
.rkrow.in{animation:rise .55s both}
.rklist .rkrow:hover{background:#f4f8fc}.rklist .rkrow.first:hover{background:#faf2e1}
@keyframes rise{from{opacity:0;transform:translateX(-14px)}to{opacity:1;transform:none}}
.rkrow .na{color:#a9b4c1}
.tb{display:flex;align-items:center;gap:10px;padding:10px 18px;border-bottom:1px solid var(--line);font-size:12px;color:var(--muted);flex-wrap:wrap}
.tb .sg{margin-left:0}
.rb{display:inline-flex;width:28px;height:28px;transition:background .4s;border-radius:50%;align-items:center;justify-content:center;background:var(--soft);color:var(--navy);font-weight:700;font-size:12.5px}
.rb.r1{background:var(--gold);color:#fff}.rb.r2{background:#8a9bb3;color:#fff}.rb.r3{background:#b08968;color:#fff}
body:not(.ready) .shell{display:none}
#gate{position:fixed;inset:0;background:var(--bg);display:flex;align-items:center;justify-content:center;z-index:100;padding:20px}
#gate[hidden]{display:none}
#gf{background:#fff;border:1px solid var(--line);border-radius:12px;padding:28px 30px;width:100%;max-width:380px;text-align:center;box-shadow:0 12px 32px rgba(18,47,89,.12)}
.glogo{width:200px;max-width:80%;margin:0 auto 18px;color:#231f20}.glogo svg{display:block;width:100%;height:auto}
#gf h2{font-size:19px;color:var(--navy)}#gf p{font-size:12.5px;color:var(--muted);margin:4px 0 16px}
#gf input{width:100%;padding:10px 12px;border:1px solid #cfd9e5;border-radius:7px;font:inherit;margin-bottom:10px}
#gf button{width:100%}.gerr{color:var(--red);font-size:12.5px;min-height:18px;margin-top:10px}
@media(max-width:1100px){.cards{grid-template-columns:repeat(2,1fr)}.hr{margin-left:0}}
@media(max-width:700px){.cards{grid-template-columns:1fr 1fr}main,.top .in{padding-left:14px;padding-right:14px}.head h1{font-size:25px}.hr{width:100%}.days{flex:1;width:auto}.crow{grid-template-columns:100px 1fr 66px;gap:8px}.tab{padding:10px 12px;font-size:14px}}
@media(prefers-reduced-motion:reduce){*{transition:none!important;animation:none!important}}
@media print{
 body{background:#fff;font-size:11px}.shell{display:none}#printArea{display:block}
 .pv{break-after:page}.pv:last-child{break-after:auto}
 .ptitle{font-size:20px;color:var(--navy);margin:0 0 2px}.psub{color:var(--muted);margin:0 0 10px;font-size:11px}
 .card{padding:8px 12px}.card .v{font-size:20px}.cards{margin-bottom:10px}.box{break-inside:avoid}.sg,.bchips{display:none}
 th,td{padding:4px 7px}.crow{padding:4px 0}
 *{-webkit-print-color-adjust:exact;print-color-adjust:exact;animation:none!important;transition:none!important}
 @page{size:A4 landscape;margin:9mm}
}
</style>
</head>
<body>
<div id="gate" hidden><form id="gf"><div class="glogo">__GATE_LOGO__</div><h2>Ek Satış Takip Paneli</h2><p>Servis Performans Yönetimi · görüntülemek için giriş yapın</p>
 <input id="gm" type="text" placeholder="Kullanıcı adı" autocapitalize="none" spellcheck="false" autocomplete="username" autofocus><input id="gp" type="password" placeholder="Şifre" autocomplete="current-password"><button class="pri" type="submit">Giriş</button><div class="gerr" id="ge"></div></form></div>
<div class="shell">
<header class="top"><div class="in">
 <div class="logo" id="logo"></div>
 <div class="acts"><a class="btn" id="dl" href="#" download="Ek_Satis_Takip.xlsx">Excel indir</a><button class="pri" onclick="window.print()">PDF</button></div>
</div></header>
<main>
 <div class="head">
  <div><div class="kick" id="kick"></div><h1 id="h1"></h1><div class="sub" id="sub"></div></div>
  <div class="hr"><div class="days" id="days"></div><div class="stamp" id="stamp"></div></div>
 </div>
 <section class="cards" id="cards"></section>
 <nav class="tabs" id="tabs">
  <button class="tab on" data-t="detay">Marka Detayı</button>
  <button class="tab" data-t="kars">Marka Karşılaştırma</button>
 </nav>
 <div id="view"></div>
 <div class="note" id="note"></div>
</main>
</div>
<div id="printArea"></div>
<script>
const PAYLOAD = __PAYLOAD__;
function start(D) {
const $ = (s, r = document) => r.querySelector(s), $$ = (s, r = document) => [...r.querySelectorAll(s)];
const M = D.months;
const reduce = window.matchMedia('(prefers-reduced-motion: reduce)').matches;
const esc = s => String(s).replace(/[&<>"]/g, c => ({'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;'}[c]));
const fN = v => v == null ? '—' : Math.round(v).toLocaleString('tr-TR');
const fP = v => v == null ? '—' : '%' + (v * 100).toLocaleString('tr-TR', {minimumFractionDigits: 1, maximumFractionDigits: 1});
const fM = v => v == null ? '—' : (v / 1e6).toLocaleString('tr-TR', {minimumFractionDigits: 2, maximumFractionDigits: 2}) + ' Mn ₺';
const FMT = {p: fP, m: fM, n: fN};

// ---------------------------------------------------------------- veri
function sumMonths(lists) {
  return M.map((_, i) => {
    let p = 0, f = null;
    lists.forEach(l => { p += l[i].p; if (l[i].f != null) f = (f || 0) + l[i].f; });
    return {p, f};
  });
}
const B = D.brands.map(b => ({ad: b.ad, items: b.items, m: sumMonths(b.items.map(i => i.m))}));
const T = sumMonths(B.map(b => b.m));
const kalemNames = [];
D.brands.forEach(b => b.items.forEach(i => { if (!kalemNames.includes(i.k)) kalemNames.push(i.k); }));
const K = kalemNames.map(k => ({k, m: sumMonths(D.brands.flatMap(b => b.items.filter(i => i.k === k).map(i => i.m)))}));

function per(m, sel) {
  const idx = sel === 'q' ? M.map((_, i) => i) : [sel];
  let p = 0, f = null, exp = 0;
  idx.forEach(i => { p += m[i].p; if (m[i].f != null) f = (f || 0) + m[i].f; exp += m[i].p * M[i].oran; });
  return {p, f, exp, gr: (f == null || !p) ? null : f / p, gd: (f == null || !exp) ? null : f / exp, kalan: Math.max(p - (f || 0), 0)};
}
const hasGid = p => p === 'q' ? M.some(m => m.durum === 'devam') : M[p].durum === 'devam';
const score = (v, p) => hasGid(p) ? v.gd : v.gr;
const cc = v => v == null ? 'n' : v >= 1 ? 'g' : v >= .9 ? 'a' : 'r';
const label = p => p === 'q' ? D.meta.donem + ' Toplam' : M[p].ad;
const focus = (() => {
  let i = M.findIndex(m => m.durum === 'devam');
  if (i < 0) i = M.map(m => m.durum).lastIndexOf('kapandi');
  return i < 0 ? 0 : i;
})();
const S = {tab: 'detay', cp: focus, brand: -1, kal: -1, sort: 'gr'};

// ---------------------------------------------------------------- KPI
function renderCards() {
  const tq = per(T, 'q'), tf = per(T, focus), fm = M[focus], g = hasGid(focus);
  const c1 = cc(g ? tf.gd : tf.gr);
  const list = [
    [`${D.meta.donem} Toplam Plan`, tq.p, 'm', 'n', `${M.length} aylık hedef`],
    ['Bugüne kadar Fiili', tq.f, 'm', 'n', tq.f == null ? 'Fiili veri girilmedi' : `Dönem planının ${fP(tq.gr)}'i`],
    [`${fm.ad} Planı`, tf.p, 'm', 'n', `Fiili: ${fM(tf.f)}`],
    [`${fm.ad} Gerçekleşme`, tf.gr, 'p', g ? 'n' : c1, g ? `Geçen iş günü ${fm.gecenGun}/${fm.toplamGun} (${fP(fm.oran)})` : (fm.durum === 'kapandi' ? 'Ay kapandı' : 'Henüz başlamadı')]
  ];
  $('#cards').innerHTML = list.map(c => `<div class="card ${c[3]}"><div class="t">${c[0]}</div><div class="v" ${c[1] != null ? `data-v="${c[1]}" data-f="${c[2]}"` : ''}>${c[1] == null ? '—' : FMT[c[2]](c[1])}</div><div class="s">${c[4]}</div></div>`).join('');
}

// ---------------------------------------------------------------- grafik
function segBtns() {
  return '<div class="sg">' + [...M.map((_, i) => i), 'q'].map(p => `<button class="${S.cp === p ? 'on' : ''}" data-cp="${p}">${p === 'q' ? D.meta.donem + ' Toplam' : M[p].ad}</button>`).join('') + '</div>';
}

// ---------------------------------------------------------------- aylık tablo
function mTable(rows, head, totLabel, totM) {
  let h1 = `<tr class="per"><th rowspan="2" style="text-align:left;vertical-align:bottom">${head}</th>`, h2 = '<tr class="sub">';
  M.forEach(m => {
    const dv = m.durum === 'devam';
    h1 += `<th colspan="3" class="${dv ? 'cur' : ''}">${m.ad} <small>${dv ? '· devam ediyor' : m.durum === 'kapandi' ? '· kapandı' : ''}</small></th>`;
    h2 += '<th>Plan</th><th>Fiili</th><th>Gerç. %</th>';
  });
  h1 += `<th colspan="3" class="ytd">${D.meta.donem} Toplam</th></tr>`;
  h2 += '<th>Plan</th><th>Fiili</th><th>Gerç. %</th></tr>';
  const cells = (m, foot) => {
    let b = '';
    m.forEach((x, i) => {
      const mo = M[i], g = x.f == null || !x.p ? null : x.f / x.p;
      const cl = mo.durum === 'devam' ? cc(g == null || !mo.oran ? null : g / mo.oran) : cc(g);
      b += `<td class="sep">${fN(x.p)}</td><td class="${x.f == null ? 'na' : ''}">${fN(x.f)}</td><td class="${foot || g == null ? (g == null ? 'na' : '') : 'pp ' + cl}">${fP(g)}</td>`;
    });
    const q = per(m, 'q');
    return b + `<td class="sep qq">${fN(q.p)}</td><td class="qq ${q.f == null ? 'na' : ''}">${fN(q.f)}</td><td class="qq ${foot || q.gr == null ? '' : 'pp ' + cc(hasGid('q') ? q.gd : q.gr)}">${fP(q.gr)}</td>`;
  };
  const body = rows.map((r, i) => `<tr style="animation-delay:${i * 45}ms"><td>${r.color ? `<i class="dot" style="background:${r.color}"></i>` : ''}${esc(r.label)}</td>${cells(r.m, false)}</tr>`).join('');
  return `<div class="scroll"><table><thead>${h1}${h2}</thead><tbody>${body}</tbody><tfoot><tr><td>${esc(totLabel)}</td>${cells(totM, true)}</tr></tfoot></table></div>`;
}
const tableBox = (title, sub, html) => `<section class="box"><div class="boxhead"><h2>${title}</h2><span class="s">${sub}</span></div>${html}</section>`;

// ---------------------------------------------------------------- sayfalar
const SORTS = [['gr', 'Gerçekleşme %'], ['f', 'Fiili'], ['p', 'Plan'], ['kalan', 'Kalan Hedef']];
let FLIPMODE = false;
function vKars() {
  const all = S.kal < 0, kn = all ? 'Tüm Kalemler' : K[S.kal].k;
  const totM = all ? T : K[S.kal].m;
  const bm = b => all ? b.m : ((b.items.find(i => i.k === kn) || {}).m || M.map(() => ({p: 0, f: null})));
  const rows = B.map(b => ({n: b.ad, v: per(bm(b), S.cp)}));
  const noF = rows.every(r => r.v.f == null);
  const key = noF ? 'p' : S.sort;
  const val = r => key === 'gr' ? r.v.gr : key === 'f' ? r.v.f : key === 'p' ? r.v.p : r.v.kalan;
  rows.sort((a, b) => (val(b) ?? -1) - (val(a) ?? -1));
  const tv = per(totM, S.cp), tc = cc(score(tv, S.cp));
  const kchip = (i, ad, m) => {
    const v = per(m, focus);
    return `<button class="bchip ${i === S.kal ? 'on' : ''}" data-kal="${i}">${esc(ad)}<small>${v.gr == null ? 'fiili yok' : fP(v.gr)}</small></button>`;
  };
  const chips = '<div class="bchips">' + kchip(-1, 'Tüm Kalemler', T) + K.map((k, i) => kchip(i, k.k, k.m)).join('') + '</div>';
  const num = (v, f) => v == null ? '—' : `<span data-v="${v}" data-f="${f}">${FMT[f](v)}</span>`;
  const list = rows.map((r, i) => {
    const v = r.v, c = cc(score(v, S.cp)), rk = i + 1;
    return `<div class="rkrow ${FLIPMODE ? '' : 'in'} ${rk === 1 && v.gr != null ? 'first' : ''}" data-k="${esc(r.n)}" style="animation-delay:${i * 60}ms">
      <div><span class="rb ${rk <= 3 ? 'r' + rk : ''}">${rk}</span></div><div class="nm">${esc(r.n)}</div>
      <div>${num(v.p, 'n')}</div><div class="${v.f == null ? 'na' : ''}">${num(v.f, 'n')}</div>
      <div>${v.gr == null ? '<span class="chip n">' + (v.p ? '—' : 'plan yok') + '</span>' : `<span class="chip ${c}" data-v="${v.gr}" data-f="p">${fP(v.gr)}</span>`}</div>
      <div>${v.f == null ? '—' : num(v.kalan, 'n')}</div></div>`;
  }).join('');
  const sortBtns = '<div class="tb"><span>Sıralama ölçütü</span><div class="sg">' + SORTS.map(x => `<button class="${key === x[0] ? 'on' : ''}" data-sort="${x[0]}">${x[1]}</button>`).join('') + '</div>' +
    (noF ? '<span>Bu dönem için fiili veri girilmedi; plana göre sıralanıyor.</span>' : '') + '</div>';
  return chips + `<section class="box"><div class="boxhead"><h2>Marka karşılaştırma · ${esc(kn)} · ${label(S.cp)}</h2><span class="s">Markalar aynı kalemde yan yana, seçilen ölçüte göre sıralı</span>${segBtns()}</div>${sortBtns}
    <div class="scroll"><div class="rkw"><div class="rkrow rkhead"><div>Sıra</div><div>Marka</div><div>Plan</div><div>Fiili</div><div>Gerç. %</div><div>Kalan Hedef</div></div>
    <div class="rklist">${list}</div>
    <div class="rkrow tot"><div></div><div class="nm">İNCİROĞLU TOPLAM</div><div>${num(tv.p, 'n')}</div><div>${num(tv.f, 'n')}</div><div>${tv.gr == null ? '—' : `<span class="chip ${tc}" data-v="${tv.gr}" data-f="p">${fP(tv.gr)}</span>`}</div><div>${tv.f == null ? '—' : num(tv.kalan, 'n')}</div></div>
    </div></div></section>`;
}
function vDetay() {
  const all = S.brand < 0;
  const b = all ? {ad: 'Tüm Markalar Toplam', m: T, items: K.map(k => ({k: k.k, m: k.m}))} : B[S.brand];
  const rows = b.items.map(i => ({n: i.k, label: i.k, m: i.m}));
  const chip = (i, ad, m) => {
    const v = per(m, focus);
    return `<button class="bchip ${i === S.brand ? 'on' : ''}" data-b="${i}">${esc(ad)}<small>${v.gr == null ? 'fiili yok' : fP(v.gr)}</small></button>`;
  };
  const chips = '<div class="bchips">' + chip(-1, 'Tüm Markalar Toplam', T) + B.map((x, i) => chip(i, x.ad, x.m)).join('') + '</div>';
  return chips + tableBox(`${esc(b.ad)} · kalem kırılımı`, 'Üç aylık plan, fiili ve gerçekleşme', mTable(rows, 'Kalem', all ? 'İNCİROĞLU TOPLAM' : b.ad + ' TOPLAM', b.m));
}
const VIEWS = {kars: vKars, detay: vDetay};

// ---------------------------------------------------------------- hareket
function settle(root, instant) {
  $$('[data-w]', root).forEach(e => e.style.width = e.dataset.w);
  $$('[data-l]', root).forEach(e => e.style.left = e.dataset.l);
  $$('[data-v]', root).forEach(el => {
    const v = +el.dataset.v, f = FMT[el.dataset.f];
    if (instant || reduce) { el.textContent = f(v); return; }
    const t0 = performance.now(), d = 950;
    const step = t => { const k = Math.min((t - t0) / d, 1); el.textContent = f(v * (1 - Math.pow(1 - k, 3))); if (k < 1) requestAnimationFrame(step); };
    requestAnimationFrame(step);
  });
}
function renderView() {
  $$('.tab').forEach(t => t.classList.toggle('on', t.dataset.t === S.tab));
  const v = $('#view');
  const wasKars = v.dataset.cur === 'kars' && S.tab === 'kars';
  const old = {};
  if (wasKars) $$('.rkrow[data-k]', v).forEach(r => old[r.dataset.k] = r.getBoundingClientRect().top);
  FLIPMODE = wasKars;
  v.innerHTML = `<div class="panel on ${wasKars ? 'noanim' : ''}">${VIEWS[S.tab]()}</div>`;
  v.dataset.cur = S.tab;
  if (wasKars && !reduce) $$('.rkrow[data-k]', v).forEach(r => {
    const o = old[r.dataset.k]; if (o == null) return;
    const dy = o - r.getBoundingClientRect().top;
    if (Math.abs(dy) < 1) return;
    r.style.transition = 'none'; r.style.transform = `translateY(${dy}px)`;
    requestAnimationFrame(() => requestAnimationFrame(() => { r.style.transition = 'transform .7s cubic-bezier(.2,.8,.2,1)'; r.style.transform = ''; }));
  });
  requestAnimationFrame(() => setTimeout(() => settle(v), 40));
}
document.addEventListener('click', e => {
  const t = e.target.closest('[data-t],[data-cp],[data-b],[data-kal],[data-sort]'); if (!t) return;
  if (t.dataset.t) S.tab = t.dataset.t;
  else if (t.dataset.cp !== undefined) S.cp = t.dataset.cp === 'q' ? 'q' : +t.dataset.cp;
  else if (t.dataset.b !== undefined) S.brand = +t.dataset.b;
  else if (t.dataset.kal !== undefined) S.kal = +t.dataset.kal;
  else if (t.dataset.sort !== undefined) S.sort = t.dataset.sort;
  renderView();
});

// ---------------------------------------------------------------- başlık / yazdırma
if (D.xlsx) {
  const bytes = Uint8Array.from(atob(D.xlsx), c => c.charCodeAt(0));
  const a = $('#dl');
  a.href = URL.createObjectURL(new Blob([bytes], {type: 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}));
  a.download = `Ek_Satis_Takip_${D.meta.donem.replace(' ', '_')}.xlsx`;
}
$('#logo').innerHTML = D.meta.logo || '<div class="wm">İnciroğlu Otomotiv</div>';
$('#kick').textContent = 'Servis Performans Yönetimi · Ek Satış';
$('#h1').textContent = `Ek Satış Takip Paneli · ${D.meta.donem}`;
$('#sub').textContent = 'Plan ve gerçekleşme · Tutarlar TL (ciro)';
const AY = ['Oca', 'Şub', 'Mar', 'Nis', 'May', 'Haz', 'Tem', 'Ağu', 'Eyl', 'Eki', 'Kas', 'Ara'];
{ const [dd, mm, yy] = D.meta.veriTarihi.split('.');
  $('#stamp').innerHTML = `<small>GÜNCELLENME<br>TARİHİ</small><b>${+dd}</b><span>${AY[+mm - 1]} ${yy}</span>`;
  $('#stamp').title = 'Son güncelleme: ' + D.meta.uretim; }
const cm = M.find(m => m.durum === 'devam');
if (cm) $('#days').innerHTML = `<div class="l"><span>${cm.ad} iş günü</span><span><b>${cm.gecenGun}</b> / ${cm.toplamGun} · ${fP(cm.oran)}</span></div><div class="bar"><i class="fb" data-w="${cm.oran * 100}%"></i><em class="cur" data-l="${cm.oran * 100}%"><span>${cm.gecenGun}. gün</span></em></div><div class="ends"><span>1</span><span>${cm.toplamGun}. iş günü</span></div>`;
else $('#days').style.display = 'none';
$('#note').innerHTML = `Gerç. % = Fiili ÷ Plan. Renkler: <span class="chip g">%100 ve üzeri</span><span class="chip a">%90–100</span><span class="chip r">%90 altı</span>` +
  (cm ? ` · Devam eden ayda renkler geçen iş günü oranına göre verilir (${cm.ad}: ${cm.gecenGun}/${cm.toplamGun} iş günü)` : '') +
  ` · Fiili girilmemiş aylar "—" gösterilir · Kaynak: Excel "Konsolide Hedef" sayfası · Güncelleme: ${D.meta.uretim}`;

window.addEventListener('beforeprint', () => {
  const st = {...S}; let h = '';
  const page = (t, body) => `<div class="pv"><h1 class="ptitle">${t}</h1><p class="psub">${D.meta.donem} Ek Satış Takip · Veri tarihi ${D.meta.veriTarihi}</p><section class="cards">${$('#cards').innerHTML}</section>${body}</div>`;
  S.cp = focus;
  S.brand = -1; h += page('Marka Detayı · Tüm Markalar Toplam', vDetay());
  B.forEach((b, i) => { S.brand = i; h += page('Marka Detayı · ' + esc(b.ad), vDetay()); });
  S.kal = -1; S.sort = 'gr';
  h += page('Marka Karşılaştırma · Tüm Kalemler', vKars());
  Object.assign(S, st);
  const pa = $('#printArea'); pa.innerHTML = h; settle(pa, true);
});
window.addEventListener('afterprint', () => { $('#printArea').innerHTML = ''; });

renderCards();
renderView();
requestAnimationFrame(() => setTimeout(() => settle(document), 40));
document.body.classList.add('ready');
}

// ---------------------------------------------------------------- açılış: düz veri ya da şifreli veri
(function () {
  if (PAYLOAD.plain) { start(PAYLOAD.plain); return; }
  const enc = PAYLOAD.enc, gate = document.getElementById('gate'), err = document.getElementById('ge');
  const b64 = x => Uint8Array.from(atob(x), c => c.charCodeAt(0));
  const te = new TextEncoder();
  async function aes(keyBytes, iv, ct) {
    const k = await crypto.subtle.importKey('raw', keyBytes, 'AES-GCM', false, ['decrypt']);
    return new Uint8Array(await crypto.subtle.decrypt({name: 'AES-GCM', iv: b64(iv)}, k, b64(ct)));
  }
  async function pbk(secret, salt, iter) {
    const km = await crypto.subtle.importKey('raw', te.encode(secret), 'PBKDF2', false, ['deriveBits']);
    return new Uint8Array(await crypto.subtle.deriveBits({name: 'PBKDF2', salt: b64(salt), iterations: iter, hash: 'SHA-256'}, km, 256));
  }
  async function decrypt(mail, pw) {
    mail = mail.trim().toLowerCase();
    const idb = await crypto.subtle.digest('SHA-256', te.encode(enc.pep + mail));
    const id = Array.from(new Uint8Array(idb)).slice(0, 8).map(x => x.toString(16).padStart(2, '0')).join('');
    const u = enc.users.find(x => x.id === id);
    if (!u) throw new Error('yok');
    const dk = await aes(await pbk(mail + ':' + pw, u.salt, enc.iter), u.iv, u.ct);
    return JSON.parse(new TextDecoder().decode(await aes(dk, enc.iv, enc.ct)));
  }
  async function attempt(mail, pw, silent) {
    try {
      const d = await decrypt(mail, pw);
      try { sessionStorage.setItem('eks_c', JSON.stringify([mail, pw])); } catch (e) {}
      gate.hidden = true; start(d);
    } catch (e) { if (!silent) { err.textContent = 'Kullanıcı adı veya şifre hatalı. Lütfen tekrar deneyin.'; } }
  }
  gate.hidden = false;
  document.getElementById('gf').addEventListener('submit', e => { e.preventDefault(); err.textContent = ''; attempt(document.getElementById('gm').value, document.getElementById('gp').value, false); });
  try { const sp = JSON.parse(sessionStorage.getItem('eks_c') || 'null'); if (sp) attempt(sp[0], sp[1], true); } catch (e) {}
})();
</script>
</body>
</html>
'''


def load_logo(folder, explicit=None):
    """Logo HTML'i döndürür. Klasörde logo.svg/png/jpg varsa onu, yoksa gömülü İnciroğlu logosunu kullanır."""
    cands = [explicit] if explicit else [os.path.join(folder, "logo" + e) for e in
                                         (".svg", ".png", ".jpg", ".jpeg", ".webp")]
    for p in cands:
        if p and os.path.exists(p):
            if p.lower().endswith(".svg"):
                with open(p, encoding="utf-8") as fh:
                    return fh.read().strip()
            mt = mimetypes.guess_type(p)[0] or "image/png"
            with open(p, "rb") as fh:
                return (f'<img src="data:{mt};base64,' + base64.b64encode(fh.read()).decode() +
                        '" alt="İnciroğlu Otomotiv" style="display:block;width:100%;height:auto">')
    return LOGO_SVG


PBKDF2_ITER = 200_000


def read_users(here):
    """kullanicilar.txt: her satır 'mail' ya da 'mail;sifre'. Şifresi olmayana otomatik şifre üretilip dosyaya yazılır."""
    p = os.path.join(here, "kullanicilar.txt")
    if os.path.exists(p):
        with open(p, encoding="utf-8-sig") as fh:
            raw = [l.strip() for l in fh if l.strip() and not l.strip().startswith("#")]
        users, lines, changed = {}, [], False
        for l in raw:
            mail, _, pw = l.partition(";")
            mail, pw = mail.strip().lower(), pw.strip()
            if not mail:
                sys.exit(f"HATA: kullanicilar.txt içinde boş kullanıcı adı: {l}")
            if not pw:
                pw, changed = gen_password(), True
            users[mail] = pw
            lines.append(f"{mail};{pw}")
        if changed:
            with open(p, "w", encoding="utf-8") as fh:
                fh.write("# mail;sifre   (bu dosya GitHub'a YÜKLENMEZ; sifreleri ilgili kisilere iletin)\n" + "\n".join(lines) + "\n")
            print("Yeni şifreler kullanicilar.txt dosyasına yazıldı.")
        return users
    return None


def gen_password():
    import secrets
    al = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    return "-".join("".join(secrets.choice(al) for _ in range(4)) for _ in range(3))


def make_payload(data, users):
    """users {mail: sifre} verilmişse veri rastgele bir anahtarla AES-GCM ile şifrelenir;
    anahtar her kullanıcının mail+şifresinden türetilen anahtarla ayrı ayrı sarılır."""
    if not users:
        return {"plain": data}
    try:
        from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    except ImportError:
        sys.exit("HATA: Şifreli çıktı için bir kez şunu çalıştırın:  pip install cryptography")
    b = lambda x: base64.b64encode(x).decode()
    dk, iv, pep = os.urandom(32), os.urandom(12), b(os.urandom(12))
    ct = AESGCM(dk).encrypt(iv, json.dumps(data, ensure_ascii=False).encode("utf-8"), None)
    out = []
    for mail, pw in users.items():
        salt, uiv = os.urandom(16), os.urandom(12)
        k = hashlib.pbkdf2_hmac("sha256", (mail + ":" + pw).encode("utf-8"), salt, PBKDF2_ITER, 32)
        uid = hashlib.sha256((pep + mail).encode("utf-8")).hexdigest()[:16]
        out.append({"id": uid, "salt": b(salt), "iv": b(uiv), "ct": b(AESGCM(k).encrypt(uiv, dk, None))})
    return {"enc": {"pep": pep, "iv": b(iv), "ct": b(ct), "iter": PBKDF2_ITER, "users": out}}


def read_password(here, explicit):
    """Eski tek-şifre yöntemi: tek kullanıcı gibi davranır (mail: genel@panel)."""
    pw = explicit or os.environ.get("EK_SATIS_SIFRE")
    p = os.path.join(here, "sifre.txt")
    if not pw and os.path.exists(p):
        with open(p, encoding="utf-8-sig") as fh:
            pw = fh.readline().strip() or None
    return {"genel@panel": pw} if pw else None


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Ek Satış Takip Paneli üretici")
    ap.add_argument("excel", nargs="?", default=os.path.join(here, "ek_satis.xlsx"))
    ap.add_argument("--tarih", help="Veri tarihi (YYYY-MM-DD); varsayılan: bugün")
    ap.add_argument("--logo", help="Logo dosyası (varsayılan: script klasöründeki logo.png/svg/jpg)")
    ap.add_argument("--sifre", help="Sayfa şifresi (ya da sifre.txt / EK_SATIS_SIFRE); verilirse veri şifrelenir")
    ap.add_argument("--cikti", default=here, help="Çıktı klasörü")
    a = ap.parse_args()
    if not os.path.exists(a.excel):
        sys.exit(f"HATA: Excel dosyası bulunamadı: {a.excel}")
    today = dt.date.fromisoformat(a.tarih) if a.tarih else dt.date.today()
    data, warns = parse(a.excel, today)
    for w in warns:
        print("UYARI:", w)
    data["meta"]["logo"] = load_logo(here, a.logo)
    os.makedirs(a.cikti, exist_ok=True)
    xlsx_path = os.path.join(a.cikti, "ek_satis_takip.xlsx")
    write_xlsx(data, xlsx_path)
    with open(xlsx_path, "rb") as fh:
        data["xlsx"] = base64.b64encode(fh.read()).decode()
    pw = read_users(here) or read_password(here, a.sifre)
    out_html = os.path.join(a.cikti, "index.html")
    with open(out_html, "w", encoding="utf-8") as f:
        f.write(HTML.replace("__PAYLOAD__", json.dumps(make_payload(data, pw), ensure_ascii=False))
                .replace("__GATE_LOGO__", data["meta"]["logo"]))
    print(f"Hazır: {out_html}  ({data['meta']['donem']}, veri tarihi {data['meta']['veriTarihi']}, "
          f"{len(data['brands'])} marka, {f'ŞİFRELİ, {len(pw)} kullanıcı' if pw else 'şifresiz'})")


if __name__ == "__main__":
    main()
